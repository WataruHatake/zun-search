"""起動中の専用Meilisearchで確認: RAG_TEST_MEILI=1 .venv/bin/pytest tests/test_meili_faiss.py -q"""

import os
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

import numpy as np
import pytest
from test_search import FakeOllama

from rag_app.config import DEFAULTS, PROJECT
from rag_app.index import Index


@pytest.mark.skipif(os.getenv("RAG_TEST_MEILI") != "1", reason="専用Meilisearchを起動して明示的に実行")
def test_search_update_repair_and_ui():
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    from rag_app import meili_faiss as backend
    from scripts.launch_meili_faiss import app_environment

    # 専用リランカーとFAISSを併用してもOpenMPの衝突でプロセスが落ちない。
    subprocess.run(
        [
            sys.executable,
            "-c",
            "import faiss, torch, numpy as np; i=faiss.IndexFlatIP(2); i.add(np.ones((1,2),dtype='float32')); assert i.search(np.ones((1,2),dtype='float32'),1)[1][0,0]==0; assert (torch.ones(2,2) @ torch.ones(2,2)).sum().item()==8",
        ],
        env=app_environment(),
        check=True,
        capture_output=True,
    )

    with TemporaryDirectory() as temporary:
        root = Path(temporary)
        folder = root / "資料"
        folder.mkdir()
        long = folder / "青空商事.txt"
        long.write_text("RAGを実装した。担当は田中。\n" * 100)
        leave = folder / "休暇.md"
        leave.write_text("有給休暇は前営業日までに申請する。")
        destination = root / "search.sqlite3"
        model = FakeOllama()
        client = backend.Meili()
        uid = None
        try:
            with backend.MeiliFaissIndex(destination) as index:
                uid = index.uid
                report = index.update(folder, model, "embed")
                assert report["search_indexes"]["faiss_vectors"] == index.stats()["vectors"]
                assert index.keyword("青空商事") and not index.keyword("火星燃料タンク")
                unsupported = MagicMock()
                unsupported.verify_local.return_value = "text-only"
                unsupported.request.return_value = {"capabilities": ["completion"]}
                with pytest.raises(RuntimeError, match="画像入力"):
                    index.update(folder, unsupported, image_model="text-only")
                assert index.keyword("青空商事")  # 更新開始前の設定エラーで既存の索引を無効にしない。
                assert not index.keyword("RAG 火星")  # 後半の条件を捨ててRAGだけで検索しない。
                assert not index.keyword(" ")
                query = model.embed("embed", ["RAG"], query=True)[0]
                expected = Index.vector(index, query, min_similarity=0)
                actual = index.vector(query, min_similarity=0)
                assert [hit[0] for hit in actual] == [hit[0] for hit in expected]
                np.testing.assert_allclose([h[1] for h in actual], [h[1] for h in expected], atol=1e-6)
                assert len(actual) == 4  # 長い資料は3箇所まで、別の資料も候補に残る。
                long_id = next(f["id"] for f in index.files() if f["relpath"] == long.name)
                leave_id = next(f["id"] for f in index.files() if f["relpath"] == leave.name)
                selected = index.vector(query, 1, min_similarity=0, file_ids={leave_id})
                assert index.chunks(ids=[selected[0][0]])[0]["file_id"] == leave_id
                assert not index.keyword("RAG", file_ids={leave_id})
                assert not index.keyword("RAG", file_ids=set())
                assert not index.keyword("RAGの担当者")  # 本文には「者」がない。
                assert index.keyword("RAGの担当者", file_id=long_id)  # 質問では「担当」までで拾える。
                assert len(index.keyword("RAG", file_id=long_id)) > 3
                selected = index.vector(query, 1, leave_id, min_similarity=0)
                assert index.chunks(ids=[selected[0][0]])[0]["file_id"] == leave_id
                with pytest.raises(RuntimeError, match="使用中"):
                    backend.MeiliFaissIndex(destination)
                calls = model.calls
                with patch.object(index, "synchronize", side_effect=AssertionError("不要な索引再作成")):
                    assert index.update(folder, model, "embed")["unchanged"] == 2
                assert model.calls == calls

                before = long.stat()
                long.write_text(long.read_text().replace("田中", "鈴木"))
                os.utime(long, ns=(before.st_atime_ns, before.st_mtime_ns))
                assert index.update(folder, model, "embed")["changed"] == 1
                assert index.keyword("鈴木") and not index.keyword("田中")
                long.rename(folder / "新しい報告.txt")
                leave.unlink()
                updated = index.update(folder, model, "embed")
                assert updated["removed"] == 2 and updated["added"] == 1
                assert not index.keyword("休暇")
                assert not index.keyword("青空商事")

                # モデル変更直後に埋め込みが失敗しても、古いモデルのFAISS索引を残さない。
                model.signature = "embedding:v2"
                model.fail_after = model.calls
                assert index.update(folder, model, "embed")["warnings"]
                assert not index.vector(query)
                model.fail_after = None
                assert index.update(folder, model, "embed")["embedded"] > 0
                with patch.object(index.meili, "task", side_effect=RuntimeError("通信中断")):
                    with pytest.raises(RuntimeError, match="通信中断"):
                        index.synchronize()
                with pytest.raises(RuntimeError, match="未完了"):
                    index.keyword("RAG")
                calls = model.calls
                index.synchronize()
                assert model.calls == calls

            with backend.MeiliFaissIndex(destination) as reopened:
                assert reopened.keyword("RAG") and reopened.vector(query)
                with pytest.raises(ValueError):
                    reopened.update(folder / "missing")
                assert reopened.stats()["files"] == 1
                reopened.synchronize()

            for number in range(12):
                (folder / f"RAG資料{number:02}.txt").write_text(f"RAGを実装した。担当は鈴木。資料{number}。")
            (folder / "RAG資料11.txt").rename(folder / "RAG資料11.MD")
            with backend.MeiliFaissIndex(destination) as reopened:
                reopened.update(folder, model, "embed")

            config = {**DEFAULTS, "folder": str(folder), "embedding_model": "embed"}
            with (
                patch.object(backend, "settings", return_value=config),
                patch.object(backend, "database_path", return_value=destination),
                patch.object(backend, "save_settings"),
                patch("rag_app.models.Ollama", return_value=model),
                patch(
                    "rag_app.service.dedicated_rerank",
                    side_effect=lambda query, texts, name: [0.9] * len(texts),
                ) as rerank,
                # AppTest未対応のpaginationを同じkey/callbackの数値入力で操作する。
                # 実際の番号ボタンと画面位置はブラウザで確認する。
                patch(
                    "streamlit.pagination",
                    side_effect=lambda pages, **kwargs: st.number_input("ページ", 1, pages, **kwargs),
                ),
            ):
                app = AppTest.from_file(str(PROJECT / "app_meili_faiss.py"), default_timeout=30).run()
                assert not app.exception and not app.error
                assert app.title[0].value == "Zun Search" and app.metric[0].value == "13"
                assert len(next(s for s in app.selectbox if s.label == "検索方式").options) == 3
                assert len(app.image) == 1  # 待機中は立ち絵だけ。
                assert not any("既存版" in b.label for b in app.button)
                assert any(b.label == "専用リランカーを準備" for b in app.button)
                assert any(b.label == "モデルをダウンロード" for b in app.button)
                assert next(c for c in app.checkbox if c.label == "画像内の文字・図・写真も読み取る").value
                assert {".png", ".jpg", ".jpeg", ".webp"} <= set(
                    app.multiselect(key="search_extensions").options
                )
                next(b for b in app.button if b.label == "資料を読み込み・更新").click().run()
                assert not app.exception and not app.error
                with backend.MeiliFaissIndex(destination) as updated:
                    assert updated.meta("image_signature").endswith(":image-v1")
                next(t for t in app.text_input if t.label == "探している資料").set_value("RAG")
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and not app.error and app.session_state.result["mode"] == "hybrid"
                rerank.assert_not_called()
                app.checkbox(key="use_reranker").check()
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and not app.error
                rerank.assert_called_once()
                assert app.session_state.result["mode"] == "dedicated"
                assert app.session_state.result["reranker_used"] == config["reranker_model"]
                app.checkbox(key="use_reranker").uncheck()
                next(b for b in app.button if b.label == "検索する").click().run()
                assert app.session_state.result["mode"] == "hybrid" and rerank.call_count == 1
                app.checkbox(key="use_reranker").check().run()
                next(s for s in app.selectbox if s.label == "検索方式").select("keyword")
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and not app.error
                assert app.session_state.result["mode"] == "keyword" and rerank.call_count == 1
                assert not any(c.key == "use_reranker" for c in app.checkbox)
                app.multiselect(key="search_extensions").set_value([".md", ".pdf"])
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and not app.error
                assert [f["relpath"] for f in app.session_state.result["files"]] == ["RAG資料11.MD"]
                assert app.session_state.result["formats"] == [".md", ".pdf"]
                app.multiselect(key="search_extensions").set_value([".pdf"])
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.session_state.result["files"]
                app.multiselect(key="search_extensions").set_value([])
                next(b for b in app.button if b.label == "検索する").click().run()
                assert len(app.session_state.result["files"]) == 13
                next(s for s in app.selectbox if s.label == "検索方式").select("hybrid").run()
                app.checkbox(key="use_reranker").check().run()
                with patch(
                    "rag_app.service.expand_query",
                    return_value={"intent": "RAG", "formats": [], "queries": []},
                ):
                    next(s for s in app.selectbox if s.label == "検索方式").select("expanded")
                    next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and not app.error
                assert app.session_state.result["mode"] == "expanded" and rerank.call_count == 1
                assert not any(c.key == "use_reranker" for c in app.checkbox)
                assert any("AIの判定：" in c.value for c in app.caption)
                assert len(app.image) == 1  # 完了後に処理中の枝豆を残さない。
                [p for p in app.number_input if p.label == "ページ"][1].set_value(2).run()
                assert [p.value for p in app.number_input if p.label == "ページ"] == [2, 2]
                assert len([b for b in app.button if b.label == "この資料に質問する →"]) == 3
                [p for p in app.number_input if p.label == "ページ"][0].set_value(1).run()
                assert [p.value for p in app.number_input if p.label == "ページ"] == [1, 1]
                [b for b in app.button if b.label == "この資料に質問する →"][-1].click().run()
                assert any(b.label == "原本を開く" for b in app.button)
                next(t for t in app.text_area if t.label == "この資料への質問").set_value("RAGの担当者")
                next(b for b in app.button if b.label == "出典付きで回答").click().run()
                assert not app.exception and not app.error
                assert app.session_state.qa["citations"]
                assert len(app.image) == 2  # 本画面と質問ウィンドウの立ち絵だけ。
                selected = app.session_state.selected
                next(b for b in app.button if b.label == "回答をリセット").click().run()
                assert not app.exception and not app.error and "qa" not in app.session_state
                assert app.session_state.selected == selected
                assert next(t for t in app.text_area if t.label == "この資料への質問").value == "RAGの担当者"
                assert not any(e.label.endswith("· 引用") for e in app.expander)
                assert next(b for b in app.button if b.label == "回答をリセット").disabled
                next(b for b in app.button if b.label == "出典付きで回答").click().run()
                assert not app.exception and not app.error and app.session_state.qa["citations"]
                del app.session_state["selected"]
                app.run()
                assert not app.text_area
                [p for p in app.number_input if p.label == "ページ"][1].set_value(2).run()
                [b for b in app.button if b.label == "この資料に質問する →"][-1].click().run()
                assert "qa" not in app.session_state
                assert not app.exception and not app.error
                del app.session_state["selected"]
                app.run()
                next(s for s in app.selectbox if s.label == "1ページの表示件数").select("すべて").run()
                assert [p.value for p in app.number_input if p.label == "ページ"] == [1, 1]
                assert len([b for b in app.button if b.label == "この資料に質問する →"]) == 13
                next(t for t in app.text_input if t.label == "探している資料").set_value(" ")
                next(b for b in app.button if b.label == "検索する").click().run()
                assert not app.exception and any("検索文は" in e.value for e in app.error)
                assert len(app.image) == 1  # エラーでも処理中の枝豆を片付ける。

            # 無関係なPPTXも混ぜ、AI検索の拡張で正解資料を候補に回収できるか確認する。
            # ここではモデル出力を固定し、言い換えの品質は別の実モデル評価で確認する。
            from pptx import Presentation

            from rag_app.service import search

            target = "研究発表.pptx"
            for name, body in (
                (target, "検索拡張生成で社内文書を探す仕組み。"),
                ("画像認識.pptx", "YOLOで金属部品の欠陥を検出する。"),
            ):
                slides = Presentation()
                slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = body
                slides.save(folder / name)
            with backend.MeiliFaissIndex(destination) as index:
                index.update(folder, model, "embed")
                assert not index.keyword("パワポ")  # 索引に手書きの別名が残っていない。
                baseline = search(index, "RAG パワポ", model, config, "hybrid")
                assert target not in baseline["candidate_files"]
                model.reply = {"queries": ["検索拡張生成 pptx"], "formats": [".pptx"]}
                with patch(
                    "rag_app.service.llm_rerank",
                    side_effect=lambda query, rows, ollama, name, intent, formats: [
                        {
                            "rerank_score": 1 if "検索拡張生成" in row["text"] else 0,
                            "rerank_reason": "本文の検索対象を確認",
                        }
                        for row in rows
                    ],
                ) as reranker:
                    result = search(index, "RAG パワポ", model, config, "expanded")
                assert not result["warnings"]
                assert target in result["candidate_files"]
                assert result["files"][0]["relpath"] == target
                assert "画像認識.pptx" in result["candidate_files"]
                assert all(f["relpath"] != "画像認識.pptx" for f in result["files"])
                assert any(row["relpath"] == target for row in reranker.call_args.args[1])
                # 他形式が上位を占めても、件数制限の前に選択形式へ絞る。
                chosen = search(index, "RAG", model, {**config, "candidate_limit": 1}, "keyword", [".md"])
                assert [f["relpath"] for f in chosen["files"]] == ["RAG資料11.MD"]
                chosen = search(
                    index, "RAG", model, {**config, "min_similarity": 0}, "hybrid", [".md", ".pptx"]
                )
                assert {Path(f["path"]).suffix.lower() for f in chosen["files"]} == {".md", ".pptx"}
        finally:
            if uid:
                client.task("DELETE", f"/indexes/{uid}")
            client.session.close()
