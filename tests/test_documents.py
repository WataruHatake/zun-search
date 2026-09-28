from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from rag_app.documents import Passage, discover, extract, read_pdf, split_passages


def test_pdf_keeps_blank_page_numbers(tmp_path):
    path = tmp_path / "sample.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(60, 700, "page one")
    pdf.showPage()
    pdf.showPage()
    pdf.drawString(60, 700, "page three")
    pdf.save()
    result = read_pdf(path)
    assert [p.location for p in result.passages] == ["1ページ", "3ページ"]
    assert result.passages[1].text == "page three"
    assert any("2ページ" in warning for warning in result.warnings)


def test_pptx_table_and_group_text(tmp_path):
    path = tmp_path / "table.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "画像認識の実装"
    table = slide.shapes.add_table(2, 2, 0, Inches(2), Inches(5), Inches(2)).table
    table.cell(0, 0).text = "方式"
    table.cell(0, 1).text = "精度"
    table.cell(1, 0).text = "YOLO"
    table.cell(1, 1).text = "98.1%"
    presentation.save(path)
    passage = extract(path).passages[0]
    assert passage.location == "スライド 1"
    assert "画像認識の実装" in passage.text
    assert "YOLO | 98.1%" in passage.text


def test_excel_coordinates_headers_and_missing_cached_formula(tmp_path):
    path = tmp_path / "cost.xlsx"
    workbook = Workbook()
    workbook.active.title = "費用"
    for row in [["項目", "費用"], ["検証", 300000], ["合計", "=SUM(B2:B2)"]]:
        workbook.active.append(row)
    workbook.save(path)
    result = extract(path)
    assert result.passages[1].location == "費用!A2:B2"
    assert "B1=費用" in result.passages[1].text
    assert "B2=300000" in result.passages[1].text
    assert any("B3" in warning and "計算結果" in warning for warning in result.warnings)


def test_discovery_ignores_hidden_and_symlinks(tmp_path):
    for relative in [
        "normal.PDF",
        "nested/test.txt",
        ".private/secret.md",
        "node_modules/ignored.txt",
        "~$lock.xlsx",
    ]:
        path = tmp_path / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("test")
    (tmp_path / "link.txt").symlink_to(tmp_path / "nested/test.txt")
    (tmp_path / "loop").symlink_to(tmp_path, target_is_directory=True)
    files, warnings = discover(tmp_path)
    assert {str(f.relative_to(tmp_path)) for f in files} == {"normal.PDF", "nested/test.txt"}
    assert not warnings


def test_chunking_preserves_text_and_location():
    body = "日本語の文書を分割します。\n" * 200
    chunks = split_passages([Passage("3ページ", body)])
    assert all(c.location == "3ページ" and len(c.text) <= 850 for c in chunks)
    restored = chunks[0].text + "".join(c.text[140:] for c in chunks[1:])
    assert restored == body.strip()


def test_generated_demo_covers_document_formats():
    folder = Path(__file__).resolve().parent.parent / "data/demo"
    files, _ = discover(folder)
    assert {path.suffix for path in files} == {".pdf", ".pptx", ".xlsx", ".md", ".txt"}
    for path in files:
        result = extract(path)
        assert result.passages, path.name


def test_excel_title_rows_cached_values_and_cells():
    path = (
        Path(__file__).resolve().parent.parent
        / "outputs/01a08545-50cb-7ce2-815d-fcc94b033ea9/Excel検索検証.xlsx"
    )
    result = extract(path)
    rows = {row.location: row.text for row in result.passages}
    assert not result.warnings
    assert "A1=青葉工業 音声認識導入費用" in rows["費用!A5:D5"]
    assert "C4=金額（円） | D4=担当" in rows["費用!A5:D5"]
    assert "C5=1800000" in rows["費用!A5:D5"] and "D5=山田" in rows["費用!A5:D5"]
    assert "C7=2040000" in rows["費用!A7:D7"]
    assert "B4=予定日 | C4=完了率" in rows["日程!A5:D5"]
    assert "B5=2026-09-25" in rows["日程!A5:D5"]
    assert "C5=0.875" in rows["日程!A5:D5"]
    assert "表示形式: 0.0%" in rows["日程!A5:D5"]
    assert "B2=007 | C2=0" in rows["在庫!A2:D2"]


def test_excel_extractor_update_preserves_other_files(tmp_path):
    import shutil
    from unittest.mock import patch

    from rag_app.index import Index

    source = (
        Path(__file__).resolve().parent.parent
        / "outputs/01a08545-50cb-7ce2-815d-fcc94b033ea9/Excel検索検証.xlsx"
    )
    folder = tmp_path / "docs"
    folder.mkdir()
    shutil.copyfile(source, folder / "検証.xlsx")
    (folder / "既存.txt").write_text("変更のない資料です。")
    with Index(tmp_path / "index.sqlite3") as index:
        index.update(folder)
        with index.db:
            index.db.execute("DELETE FROM meta WHERE key='extractor_xlsx'")
        with patch("rag_app.index.extract", wraps=extract) as reader:
            report = index.update(folder)
        assert [call.args[0].name for call in reader.call_args_list] == ["検証.xlsx"]
        assert report["changed"] == 1 and report["unchanged"] == 1
        with patch("rag_app.index.extract", side_effect=AssertionError("再抽出不要")):
            assert index.update(folder)["unchanged"] == 2
