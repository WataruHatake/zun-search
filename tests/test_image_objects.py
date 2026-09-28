"""画像とOffice図形の読み取り、出典、キャッシュを検証する。"""

import base64
import json
from io import BytesIO
from unittest.mock import MagicMock
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from rag_app.documents import DRAWING_NS, chart_text, discover, drawing_text, extract
from rag_app.index import Index
from rag_app.models import ModelError, Ollama


def test_excel_group_text_image_anchor_and_external_link(tmp_path):
    picture = tmp_path / "image.png"
    Image.new("RGB", (40, 40), "green").save(picture)
    path = tmp_path / "objects.xlsx"
    book = Workbook()
    book.active.title = "設計"
    book.active["A1"] = "工程設計"
    book.active.add_image(SheetImage(picture), "M23")
    book.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    xdr = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    drawing = ET.fromstring(entries["xl/drawings/drawing1.xml"])
    anchor = drawing[0]
    pic = anchor.find(f"{{{xdr}}}pic")
    anchor.remove(pic)
    group = ET.SubElement(anchor, f"{{{xdr}}}grpSp")
    group.append(
        ET.fromstring(f'''<xdr:sp xmlns:xdr="{xdr}" xmlns:a="{DRAWING_NS}">
      <xdr:nvSpPr><xdr:cNvPr id="2" name="手順"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>
      <xdr:spPr/><xdr:txBody><a:bodyPr/><a:lstStyle/>
      <a:p><a:r><a:t>受入</a:t></a:r><a:r><a:t>検査</a:t></a:r></a:p>
      <a:p><a:r><a:t>担当：佐藤</a:t></a:r></a:p></xdr:txBody></xdr:sp>''')
    )
    group.append(pic)
    entries["xl/drawings/drawing1.xml"] = ET.tostring(drawing)
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    reader = MagicMock(return_value="画像内の文字（AI読み取り）: 検査合格")
    result = extract(path, reader)
    assert any(p.text == "受入検査\n担当：佐藤" and "設計!M23付近" in p.location for p in result.passages)
    assert any("検査合格" in p.text and "画像・AI読み取り" in p.location for p in result.passages)
    reader.assert_called_once_with(picture.read_bytes())
    assert not result.retry
    # 外部URLは取得せず警告し、図形の文字は残す。
    relpath = "xl/drawings/_rels/drawing1.xml.rels"
    rels = ET.fromstring(entries[relpath])
    rels[0].set("TargetMode", "External")
    rels[0].set("Target", "https://example.invalid/private.png")
    entries[relpath] = ET.tostring(rels)
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    reader.reset_mock()
    result = extract(path, reader)
    reader.assert_not_called()
    assert any("外部リンク" in w for w in result.warnings)
    assert any("担当：佐藤" in p.text for p in result.passages)


def test_powerpoint_images_charts_and_partial_failure(tmp_path):
    picture = tmp_path / "image.png"
    Image.new("RGB", (40, 40), "green").save(picture)
    path = tmp_path / "objects.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "受入検査の結果"
    group.shapes.add_picture(str(picture), 0, Inches(1))
    data = CategoryChartData()
    data.categories = ["4月", "5月"]
    data.add_series("検査件数", [12, 18])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(5), Inches(3), data)
    deck.save(path)
    result = extract(path, lambda data: "画像内の文字（AI読み取り）: 検査合格")
    text = "\n".join(p.text for p in result.passages)
    assert "受入検査の結果" in text and "検査合格" in text
    assert "系列 1: 検査件数" in text and "項目=4月 | 値=12" in text
    assert all("スライド 1" in p.location for p in result.passages)
    failed = extract(path, MagicMock(side_effect=ModelError("接続中断")))
    assert failed.retry and any("受入検査の結果" in p.text for p in failed.passages)
    assert any("画像読み取り失敗" in warning for warning in failed.warnings)


def test_image_cache_retry_and_model_change(tmp_path):
    folder = tmp_path / "docs"
    folder.mkdir()
    Image.new("RGB", (40, 40), "green").save(folder / "one.PNG")
    (folder / "two.png").write_bytes((folder / "one.PNG").read_bytes())
    (folder / "unchanged.txt").write_text("変更しない資料")
    assert len(discover(folder)[0]) == 3
    client = MagicMock(spec=Ollama)
    client.verify_local.return_value = "vision@v1"
    client.request.return_value = {"capabilities": ["vision"]}
    client.read_image.return_value = "画像内の文字（AI読み取り）: 検査合格"
    with Index(tmp_path / "index.sqlite3") as index:
        assert index.update(folder, client, image_model="vision")["added"] == 3
        assert client.read_image.call_count == 1
        assert index.update(folder, client, image_model="vision")["unchanged"] == 3
    with Index(tmp_path / "index.sqlite3") as index:
        (folder / "one.PNG").rename(folder / "renamed.PNG")
        index.update(folder, client, image_model="vision")
        assert client.read_image.call_count == 1  # 再起動・移動後も再利用。
        client.verify_local.return_value = "vision@v2"
        assert index.update(folder, client, image_model="vision")["changed"] == 2
        assert client.read_image.call_count == 2
        Image.new("RGB", (40, 40), "blue").save(folder / "new.png")
        client.read_image.side_effect = ModelError("中断")
        assert index.update(folder, client, image_model="vision")["errors"] == 1
        client.read_image.side_effect = None
        assert index.update(folder, client, image_model="vision")["changed"] == 1
        assert index.update(folder, client, image_model="vision")["unchanged"] == 4


def test_image_api_payload_and_invalid_response():
    stream = BytesIO()
    Image.new("RGBA", (3000, 1000), (0, 128, 0, 128)).save(stream, format="PNG")
    client = Ollama()
    client.verify_local = MagicMock(return_value="vision@v1")
    client.request = MagicMock(
        return_value={"message": {"content": json.dumps({"text": "検査合格", "description": "緑色の図"})}}
    )
    try:
        result = client.read_image("vision", stream.getvalue())
        assert "画像内の文字（AI読み取り）" in result and "図・写真の説明（AI推定）" in result
        payload = client.request.call_args.args[1]
        image = Image.open(BytesIO(base64.b64decode(payload["messages"][1]["images"][0])))
        assert image.width == 2048 and image.mode == "RGB"
        assert payload["keep_alive"] == "2m"
        client.request.return_value = {"message": {"content": '{"text": null, "description": ""}'}}
        with pytest.raises(ModelError, match="形式"):
            client.read_image("vision", stream.getvalue())
        client.request.return_value = {
            "done_reason": "length",
            "message": {"content": '{"text": "途中", "description": ""}'},
        }
        with pytest.raises(ModelError, match="出力上限"):
            client.read_image("vision", stream.getvalue())
    finally:
        client.session.close()


def test_drawing_breaks_chart_cache_and_smartart(tmp_path):
    assert (
        drawing_text(
            ET.fromstring(
                f'<a:p xmlns:a="{DRAWING_NS}"><a:r><a:t>検査</a:t></a:r><a:br/><a:r><a:t>完了</a:t></a:r></a:p>'
            )
        )
        == "検査\n完了"
    )
    data = b"""<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
    <c:ser><c:val><c:numLit><c:formatCode>0.0%</c:formatCode><c:pt idx="0"><c:v>0.875</c:v></c:pt></c:numLit></c:val></c:ser>
    <c:ser/></c:chartSpace>"""
    text = chart_text(data)
    assert "0.875 (表示形式: 0.0%)" in text and "保存済みのグラフデータなし" in text
    path = tmp_path / "smartart.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    d = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    slide = ET.fromstring(entries["ppt/slides/slide1.xml"])
    slide.find(f"{{{p}}}cSld/{{{p}}}spTree").append(
        ET.fromstring(f'''<p:graphicFrame xmlns:p="{p}" xmlns:a="{DRAWING_NS}" xmlns:dgm="{d}" xmlns:r="{r}">
      <p:nvGraphicFramePr><p:cNvPr id="99" name="工程SmartArt"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
      <p:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="1000000"/></p:xfrm>
      <a:graphic><a:graphicData uri="{d}"><dgm:relIds r:dm="rIdDiagram"/></a:graphicData></a:graphic></p:graphicFrame>''')
    )
    entries["ppt/slides/slide1.xml"] = ET.tostring(slide)
    relpath = "ppt/slides/_rels/slide1.xml.rels"
    rels = ET.fromstring(entries[relpath])
    ET.SubElement(
        rels,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        Id="rIdDiagram",
        Type=r + "/diagramData",
        Target="../diagrams/data1.xml",
    )
    entries[relpath] = ET.tostring(rels)
    entries["ppt/diagrams/data1.xml"] = (
        f'<dgm:dataModel xmlns:dgm="{d}" xmlns:a="{DRAWING_NS}"><dgm:ptLst><dgm:pt><dgm:t><a:p><a:r><a:t>受付から検査へ</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>'.encode()
    )
    types = ET.fromstring(entries["[Content_Types].xml"])
    ET.SubElement(
        types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        PartName="/ppt/diagrams/data1.xml",
        ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
    )
    entries["[Content_Types].xml"] = ET.tostring(types)
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    result = extract(path)
    assert any("受付から検査へ" in passage.text for passage in result.passages)
    assert any("SmartArt" in warning for warning in result.warnings)
