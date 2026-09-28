"""元のページ・スライド・セル位置を保って資料を読み取る。"""

import hashlib
import os
import posixpath
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
EXTENSIONS = {".pdf", ".pptx", ".xlsx", ".md", ".txt"} | IMAGE_EXTENSIONS
SKIP_DIRS = {"node_modules", "venv", "__pycache__", "site-packages"}
EXTRACTOR_VERSION = "3"
XLSX_EXTRACTOR_VERSION = "4"
PPTX_EXTRACTOR_VERSION = "2"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


@dataclass
class Passage:
    location: str
    text: str


@dataclass
class Extraction:
    passages: list[Passage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    retry: bool = False


def add_image(result, data, location, image_reader):
    if image_reader is None:
        result.warnings.append(f"{location}: 画像読み取りが無効です。")
        return
    try:
        text = image_reader(data)
        if text:
            result.passages.append(Passage(location + "（画像・AI読み取り）", text))
        else:
            result.warnings.append(f"{location}: 画像から検索できる内容を取得できませんでした。")
    except Exception as exc:
        result.warnings.append(f"{location}: 画像読み取り失敗: {exc}")
        result.retry = True


def drawing_text(element):
    return "\n".join(
        "".join(
            "\n" if node.tag == f"{{{DRAWING_NS}}}br" else node.text or ""
            for node in paragraph.iter()
            if node.tag in {f"{{{DRAWING_NS}}}t", f"{{{DRAWING_NS}}}br"}
        )
        for paragraph in element.iter(f"{{{DRAWING_NS}}}p")
    ).strip()


def chart_text(data):
    root = ET.fromstring(data)
    lines = [drawing_text(root)]
    for number, series in enumerate(root.iter(f"{{{CHART_NS}}}ser"), 1):
        title = series.find(f"{{{CHART_NS}}}tx")
        name = " ".join(n.text or "" for n in title.iter(f"{{{CHART_NS}}}v")) if title is not None else ""
        lines.append(f"系列 {number}: {name}")
        # 保存済みの点番号で項目と値を対応付ける。数式・外部データは再計算しない。
        points = {}
        for tag, label in (
            ("cat", "項目"),
            ("val", "値"),
            ("xVal", "X"),
            ("yVal", "Y"),
            ("bubbleSize", "大きさ"),
        ):
            axis = series.find(f"{{{CHART_NS}}}{tag}")
            if axis is None:
                continue
            code = axis.findtext(f".//{{{CHART_NS}}}formatCode", "General")
            suffix = f" (表示形式: {code})" if code != "General" else ""
            for point in axis.iter(f"{{{CHART_NS}}}pt"):
                points.setdefault(point.get("idx"), []).append(
                    f"{label}={point.findtext(f'{{{CHART_NS}}}v', '')}{suffix}"
                )
        lines.extend(" | ".join(values) for values in points.values())
        if not points:
            lines.append("[保存済みのグラフデータなし]")
    return "\n".join(filter(None, lines))


def relationships(archive, part):
    folder, name = posixpath.split(part)
    path = posixpath.join(folder, "_rels", name + ".rels")
    if path not in archive.namelist():
        return {}
    return {
        rel.get("Id"): posixpath.normpath(posixpath.join(folder, rel.get("Target", ""))).lstrip("/")
        for rel in ET.fromstring(archive.read(path))
        if rel.get("TargetMode") != "External"
    }


def discover(folder: Path) -> tuple[list[Path], list[str]]:
    folder = folder.expanduser().resolve()
    if not folder.is_dir():
        raise ValueError(f"対象フォルダが見つかりません: {folder}")
    files, warnings = [], []
    for base, dirs, names in os.walk(folder, followlinks=False, onerror=lambda e: warnings.append(str(e))):
        dirs[:] = sorted(
            d for d in dirs if not d.startswith(".") and d not in SKIP_DIRS and not Path(base, d).is_symlink()
        )
        for name in sorted(names):
            path = Path(base, name)
            if name.startswith((".", "~$")) or path.is_symlink():
                continue
            if path.suffix.lower() in EXTENSIONS and path.is_file():
                files.append(path)
    return sorted(files), warnings


def fingerprint(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def read_pdf(path: Path) -> Extraction:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer

    result = Extraction()
    # 改ページ文字では分割しない。空白ページも数えて原本のページ番号を保つ。
    for page_number, layout in enumerate(extract_pages(path), 1):
        text = "\n".join(item.get_text() for item in layout if isinstance(item, LTTextContainer)).strip()
        if text:
            result.passages.append(Passage(f"{page_number}ページ", text))
            if text.count("(cid:") > 3 or text.count("\ufffd") > 3:
                result.warnings.append(f"{page_number}ページ: 文字化けの可能性があります。")
        else:
            result.warnings.append(
                f"{page_number}ページ: 文字を取得できません。画像・スキャンPDFはOCRが必要です。"
            )
    return result


def read_pptx(path: Path, image_reader=None) -> Extraction:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    result = Extraction()

    def collect(shapes, location):
        parts = []
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                parts.extend(collect(shape.shapes, location))
                continue
            elif shape.has_table:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            elif shape.has_text_frame:
                parts.append(shape.text_frame.text)
            elif shape.has_chart:
                parts.append(chart_text(shape.chart.part.blob))
            object_location = f"{location} / {shape.name}"
            for blip in shape.element.iter(f"{{{DRAWING_NS}}}blip"):
                image_id = blip.get(f"{{{REL_NS}}}embed")
                if image_id:
                    add_image(result, shape.part.related_part(image_id).blob, object_location, image_reader)
                else:
                    result.warnings.append(f"{object_location}: 外部リンク画像は読み取りません。")
            diagram_ns = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
            for diagram in shape.element.iter(f"{{{diagram_ns}}}relIds"):
                diagram_id = diagram.get(f"{{{REL_NS}}}dm")
                if diagram_id:
                    parts.append(drawing_text(ET.fromstring(shape.part.related_part(diagram_id).blob)))
                    result.warnings.append(
                        f"{object_location}: SmartArtの文字を抽出しました。配置や矢印の関係は解析しません。"
                    )
            if shape.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                result.warnings.append(f"{object_location}: 埋め込みオブジェクトの本体は読み取り対象外です。")
        return parts

    for number, slide in enumerate(Presentation(path).slides, 1):
        location = f"スライド {number}"
        text = "\n".join(collect(slide.shapes, location)).strip()
        if text:
            result.passages.append(Passage(location, text))
        else:
            result.warnings.append(f"{location}: 文字を取得できません。")
    result.warnings = list(dict.fromkeys(result.warnings))
    return result


def read_xlsx(path: Path, image_reader=None) -> Extraction:
    from openpyxl import load_workbook

    result = Extraction()
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = None
    try:
        formulas = load_workbook(path, data_only=False, read_only=True)
        for sheet in values:
            formula_sheet = formulas[sheet.title]
            header = ""
            header_width = 0
            if sheet.sheet_state != "visible":
                result.warnings.append(f"{sheet.title}: 非表示シートを含めて読み込みました。")
            for value_row, formula_row in zip(sheet.iter_rows(), formula_sheet.iter_rows()):
                cells = []
                for value_cell, formula_cell in zip(value_row, formula_row):
                    value = value_cell.value
                    if formula_cell.data_type == "f" and value is None:
                        result.warnings.append(
                            f"{sheet.title}!{formula_cell.coordinate}: 数式の保存済み計算結果がありません。"
                        )
                        value = f"[計算結果なし: {formula_cell.value}]"
                    if value is not None:
                        coordinate = formula_cell.coordinate
                        if type(value) in (int, float) and value_cell.number_format != "General":
                            value = f"{value} (表示形式: {value_cell.number_format})"
                        cells.append(f"{coordinate}={value}")
                if not cells:
                    continue
                text = " | ".join(cells)
                first, last = cells[0].split("=", 1)[0], cells[-1].split("=", 1)[0]
                location = f"{sheet.title}!{first}:{last}"
                # ponytail: 先頭の複数セル行を列見出しとする。多段見出し・複数表は別途構造解析が必要。
                if not header:
                    header = text[:1000]
                    header_width = len(cells)
                    body = text
                else:
                    body = f"見出し: {header}\n{text}"
                    if header_width == 1 and len(cells) > 1:
                        header = header + "\n" + text[:1000]
                        header_width = len(cells)
                result.passages.append(Passage(location, body))
    finally:
        values.close()
        if formulas is not None:
            formulas.close()
    read_xlsx_objects(path, result, image_reader)
    return result


def read_xlsx_objects(path, result, image_reader):
    from openpyxl.utils.cell import get_column_letter

    sheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    xdr = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_paths = relationships(archive, "xl/workbook.xml")
        for sheet in workbook.iter(f"{{{sheet_ns}}}sheet"):
            part = sheet_paths[sheet.get(f"{{{REL_NS}}}id")]
            sheet_xml = ET.fromstring(archive.read(part))
            related = relationships(archive, part)
            if sheet_xml.find(f".//{{{sheet_ns}}}oleObject") is not None:
                result.warnings.append(
                    f"{sheet.get('name')}: 埋め込みオブジェクトの本体は読み取り対象外です。"
                )
            if sheet_xml.find(f"{{{sheet_ns}}}legacyDrawing") is not None:
                result.warnings.append(
                    f"{sheet.get('name')}: 旧形式の図形・コメント（VML）は読み取り対象外です。"
                )
            for drawing in sheet_xml.findall(f"{{{sheet_ns}}}drawing"):
                try:
                    drawing_path = related[drawing.get(f"{{{REL_NS}}}id")]
                    drawing_xml = ET.fromstring(archive.read(drawing_path))
                    objects = relationships(archive, drawing_path)
                    for number, anchor in enumerate(drawing_xml, 1):
                        origin = anchor.find(f"{{{xdr}}}from")
                        cell = ""
                        if origin is not None:
                            col = int(origin.findtext(f"{{{xdr}}}col", "0")) + 1
                            row = int(origin.findtext(f"{{{xdr}}}row", "0")) + 1
                            cell = f"!{get_column_letter(col)}{row}付近"
                        location = f"{sheet.get('name')}{cell} / 図形 {number}"
                        text = drawing_text(anchor)
                        if text:
                            result.passages.append(Passage(location, text))
                        for blip in anchor.iter(f"{{{DRAWING_NS}}}blip"):
                            target = objects.get(blip.get(f"{{{REL_NS}}}embed"))
                            if target:
                                add_image(result, archive.read(target), location, image_reader)
                            else:
                                result.warnings.append(f"{location}: 外部リンク画像は読み取りません。")
                        for chart in anchor.iter(f"{{{CHART_NS}}}chart"):
                            target = objects.get(chart.get(f"{{{REL_NS}}}id"))
                            if target:
                                text = chart_text(archive.read(target))
                                if text:
                                    result.passages.append(Passage(location + " / グラフの保存値", text))
                except (KeyError, ValueError, OSError, ET.ParseError) as exc:
                    result.warnings.append(f"{sheet.get('name')}: 図形読み取り失敗: {exc}")
                    result.retry = True


def read_text(path: Path) -> Extraction:
    result = Extraction()
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    for start in range(0, len(lines), 15):
        text = "\n".join(lines[start : start + 15]).strip()
        if text:
            result.passages.append(Passage(f"{start + 1}–{min(start + 15, len(lines))}行", text))
    return result


def extract(path: Path, image_reader=None) -> Extraction:
    readers = {".pdf": read_pdf, ".pptx": read_pptx, ".xlsx": read_xlsx, ".txt": read_text, ".md": read_text}
    extension = path.suffix.lower()
    if extension in IMAGE_EXTENSIONS:
        result = Extraction()
        add_image(result, path.read_bytes(), "画像", image_reader)
    elif extension in {".xlsx", ".pptx"}:
        result = readers[extension](path, image_reader)
    else:
        result = readers[extension](path)
    if not result.passages:
        result.warnings.append("本文が空です。ファイル名・保存場所のみ検索できます。")
    return result


def split_passages(passages: list[Passage], size: int = 850, overlap: int = 140) -> list[Passage]:
    if not 0 <= overlap < size:
        raise ValueError("チャンクの重なりはサイズより小さくしてください。")
    chunks = []
    for passage in passages:
        text = passage.text.replace("\x00", "").strip()
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                boundary = max(
                    text.rfind("\n", start + size // 2, end), text.rfind("。", start + size // 2, end)
                )
                if boundary >= 0:
                    end = boundary + 1
            chunks.append(Passage(passage.location, text[start:end]))
            if end == len(text):
                break
            start = end - overlap
    return chunks
