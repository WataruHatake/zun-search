"""公開・検証用の架空資料を作る。個人のDocumentsは変更しない。"""

import json
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
DEMO = ROOT / "data" / "demo"


def pdf(path, pages):
    path.parent.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    document = canvas.Canvas(str(path), pagesize=A4)
    document.setTitle(path.stem)
    width, height = A4
    for number, (title, lines) in enumerate(pages, 1):
        document.setFillColor(colors.HexColor("#247366"))
        document.rect(0, height - 12, width, 12, fill=True, stroke=False)
        document.setFillColor(colors.HexColor("#63758b"))
        document.setFont("HeiseiKakuGo-W5", 9)
        document.drawString(48, height - 55, "LOCAL RAG / DEMO DOCUMENT")
        document.setFillColor(colors.HexColor("#213047"))
        document.setFont("HeiseiKakuGo-W5", 21)
        document.drawString(48, height - 98, title)
        document.setFont("HeiseiKakuGo-W5", 12)
        y = height - 150
        for line in lines:
            for offset in range(0, len(line), 36):
                document.drawString(48, y, line[offset : offset + 36])
                y -= 22
            y -= 12
        document.setFillColor(colors.HexColor("#63758b"))
        document.setFont("HeiseiKakuGo-W5", 9)
        document.drawString(48, 42, "動作確認用の架空資料です。")
        document.drawRightString(width - 48, 42, str(number))
        document.showPage()
    document.save()


def slides(path, contents):
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    for title, lines in contents:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.clear()
        for number, line in enumerate(lines):
            paragraph = body.paragraphs[0] if number == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(22)
    presentation.save(path)


def main():
    pdf(
        DEMO / "青空商事" / "RAG_実装報告.pdf",
        [
            (
                "青空商事 RAG実装報告",
                [
                    "2026年8月、青空商事の社内文書検索にRAGを実装した。",
                    "検索にはFAISSと日本語全文検索を組み合わせた。",
                    "回答の生成にはOllamaのローカルLLMを使用した。",
                    "プロジェクト責任者は佐藤、実装担当は田中。",
                ],
            ),
            (
                "評価結果と残された課題",
                [
                    "問い合わせ対応時間は平均20分から8分へ短縮した。",
                    "評価用の検索問題50件のうち、46件で正解資料が上位5件に入った。",
                    "画像のみのPDFは読み取れず、OCRの導入が今後の課題である。",
                    "本資料は導入検討ではなく、実装と評価を完了した報告である。",
                ],
            ),
        ],
    )
    slides(
        DEMO / "青空商事" / "RAG_導入検討.pptx",
        [
            (
                "青空商事 RAG導入検討",
                ["2026年6月の検討資料。実装はまだ行っていない。", "目的は社内文書を自然文で検索すること。"],
            ),
            (
                "検討する構成",
                [
                    "FAISSと全文検索の組み合わせを候補にする。",
                    "セキュリティ上の理由からローカルLLMを検討する。",
                ],
            ),
        ],
    )
    slides(
        DEMO / "研究" / "画像認識_評価.pptx",
        [
            (
                "製造ラインの画像認識研究",
                ["YOLOを使った部品の欠陥検出を実装した。", "対象は金属部品の表面傷である。"],
            ),
            (
                "実験と評価",
                ["テスト画像1000枚で正解率98.1%を確認した。", "照明条件が暗い場合は再現率が低下した。"],
            ),
        ],
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "費用と担当"
    for row in [
        ["プロジェクト", "項目", "金額（円）", "担当"],
        ["青空商事 RAG", "開発費", 1200000, "田中"],
        ["青空商事 RAG", "検証費", 300000, "佐藤"],
        ["青空商事 RAG", "合計", "=SUM(C2:C3)", "佐藤"],
    ]:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 15
    workbook.save(DEMO / "青空商事" / "PJ費用.xlsx")
    texts = {
        "技術/RAG_基礎.md": "# RAGの基礎\nRAGは検索拡張生成であり、検索した資料の情報をLLMへ渡して回答を作る方法である。\nこの記事は技術の紹介であり、個別の会社への導入や実装結果は含まない。",
        "青空商事/ネットワーク改修.txt": "青空商事のネットワーク改修。2026年7月に社内VPNの設定を変更した。DNSの名前解決に関する障害はキャッシュの更新で解決した。RAGの導入とは別の案件である。",
        "白波工業/音声認識_導入.md": "# 白波工業 音声認識導入\nWhisperを使って工場内の音声を文字起こしするアプリを実装した。雑音が大きい環境ではマイクの設置位置が重要だった。担当は鈴木。",
        "総務/休暇手続き.md": "# 休暇申請\n有給休暇の申請は、前営業日までに勤怠システムで行う。上長の承認後にカレンダーへ登録する。問い合わせ先は総務の山本。",
    }
    for relative, content in texts.items():
        path = DEMO / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
    pdf(
        DEMO / "技術" / "Git_基本操作.pdf",
        [
            (
                "Gitの基本操作",
                [
                    "git statusで作業ツリーの変更を確認する。",
                    "git addで変更をステージし、git commitで履歴を記録する。",
                    "ブランチはgit switch -cで作成する。",
                ],
            )
        ],
    )
    problems = [
        ("青空商事でRAGを実装し評価した報告", ["青空商事/RAG_実装報告.pdf"]),
        ("青空商事のRAG導入を検討している資料", ["青空商事/RAG_導入検討.pptx"]),
        ("問い合わせ対応時間を短縮した結果", ["青空商事/RAG_実装報告.pdf"]),
        ("青空商事の開発費と担当者", ["青空商事/PJ費用.xlsx"]),
        ("YOLOで金属部品の欠陥検出を実装した研究", ["研究/画像認識_評価.pptx"]),
        ("Whisperで工場の音声を文字起こしした資料", ["白波工業/音声認識_導入.md"]),
        ("VPNのDNS障害を解決した案件", ["青空商事/ネットワーク改修.txt"]),
        ("Gitのブランチを作るコマンド", ["技術/Git_基本操作.pdf"]),
        ("有給休暇の申請手順", ["総務/休暇手続き.md"]),
        ("火星探査機の燃料タンクの設計", []),
    ]
    dataset = [{"query": q, "relevant_files": files, "split": "demo"} for q, files in problems]
    (ROOT / "data" / "evaluation_demo.json").write_text(json.dumps(dataset, ensure_ascii=False, indent=2))
    print(f"サンプル資料: {DEMO}")


if __name__ == "__main__":
    main()
